"""Peer comparison Excel workbook generator with 3-tier conditional color scales.

Sprint 3, Day 19: Peer Comparison Excel Workbook
Generates peer_comparison.xlsx with 11 intra-peer group sheets and a master 92-company
Summary sheet, embedding 20 financial metrics, benchmark gaps, bold benchmark headers,
and openpyxl ColorScaleRule conditional formatting.
"""

from __future__ import annotations

import logging
import shutil
import sqlite3
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.screener.engine import (
    DEFAULT_DB_PATH,
    load_screener_universe,
    run_all_presets,
)
from src.screener.ranking import add_composite_scores

logger = logging.getLogger(__name__)

DEFAULT_OUTPUT_FILE = Path("reports/peer_comparison.xlsx")

# 20 Metrics: (db_metric_key, display_label, lower_is_better)
PEER_WORKBOOK_METRICS: list[tuple[str, str, bool]] = [
    ("ROE", "ROE", False),
    ("ROCE", "ROCE", False),
    ("NPM", "NPM", False),
    ("OPM", "OPM", False),
    ("D/E", "D/E", True),
    ("ICR", "ICR", False),
    ("Revenue CAGR 3yr", "Revenue CAGR 3yr", False),
    ("PAT CAGR 3yr", "PAT CAGR 3yr", False),
    ("EPS CAGR 3yr", "EPS CAGR 3yr", False),
    ("FCF", "FCF", False),
    ("CFO Quality Score", "CFO Quality Score", False),
    ("Asset Turnover", "Asset Turnover", False),
    ("CapEx Intensity", "CapEx Intensity", True),
    ("Net Profit Margin", "Net Profit Margin", False),
    ("Composite Score", "Composite Score", False),
    ("PE ratio", "PE", True),
    ("PB ratio", "PB", True),
    ("EV/EBITDA", "EV/EBITDA", True),
    ("Dividend Yield", "Dividend Yield", False),
    ("Book Value per Share", "Book Value per Share", False),
]

# Color codes
COLOR_GREEN = "63BE7B"  # Top performer
COLOR_YELLOW = "FFEB84"  # Middle
COLOR_RED = "F8696B"  # Bottom performer


def _style_header_cell(
    cell, is_benchmark: bool = False, is_metric_col: bool = False
) -> None:
    """Apply styling to header row cell."""
    if is_metric_col:
        cell.fill = PatternFill(
            start_color="1E3A8A", end_color="1E3A8A", fill_type="solid"
        )
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="left", vertical="center")
    elif is_benchmark:
        cell.fill = PatternFill(
            start_color="1E40AF", end_color="1E40AF", fill_type="solid"
        )
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    else:
        cell.fill = PatternFill(
            start_color="3B82F6", end_color="3B82F6", fill_type="solid"
        )
        cell.font = Font(name="Calibri", size=11, bold=False, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _style_section_banner(ws, row_idx: int, title: str, max_col: int) -> None:
    """Format a full-width section header banner."""
    cell = ws.cell(row=row_idx, column=1, value=title)
    cell.fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    cell.font = Font(name="Calibri", size=10, bold=True, color="F8FAFC")
    cell.alignment = Alignment(horizontal="left", vertical="center")

    for col in range(2, max_col + 1):
        c = ws.cell(row=row_idx, column=col)
        c.fill = PatternFill(
            start_color="0F172A", end_color="0F172A", fill_type="solid"
        )


def generate_peer_comparison_workbook(
    output_path: Path | str = DEFAULT_OUTPUT_FILE,
    db_path: Path | str = DEFAULT_DB_PATH,
) -> Path:
    """Build the master peer comparison Excel workbook with 12 sheets.

    Sheets:
      1. Summary (all 92 companies ranked with preset membership flags)
      2-12. 11 Peer Group sheets with 20 metrics, ColorScaleRule conditional formatting,
            benchmark gap rows, and bold benchmark company headers.

    Args:
        output_path: Destination path for peer_comparison.xlsx.
        db_path: Path to SQLite database.

    Returns:
        Path to the saved Excel workbook.
    """
    out_file = Path(output_path)
    out_file.parent.mkdir(parents=True, exist_ok=True)

    resolved_db = Path(db_path).resolve()
    if not resolved_db.exists():
        raise FileNotFoundError(f"Database not found at: {resolved_db}")

    # Load data sources
    universe_df = load_screener_universe(db_path=resolved_db)
    if "composite_ranking_score" not in universe_df.columns:
        universe_df = add_composite_scores(universe_df)

    all_presets = run_all_presets(df=universe_df, db_path=resolved_db)

    conn = sqlite3.connect(f"file:{resolved_db.as_posix()}?mode=ro", uri=True)
    try:
        peer_groups_df = pd.read_sql_query(
            "SELECT peer_group_name, company_id, is_benchmark FROM peer_groups ORDER BY peer_group_name, is_benchmark DESC, company_id ASC;",
            conn,
        )
        peer_percentiles_df = pd.read_sql_query(
            "SELECT company_id, peer_group_name, metric_name, metric_value, percentile_rank, classification, benchmark_gap_pct FROM peer_percentiles;",
            conn,
        )
        companies_df = pd.read_sql_query(
            "SELECT id, company_name FROM companies;",
            conn,
        )
    finally:
        conn.close()

    comp_name_map = dict(
        zip(companies_df["id"], companies_df["company_name"], strict=False)
    )
    peer_group_map = dict(
        zip(
            peer_groups_df["company_id"],
            peer_groups_df["peer_group_name"],
            strict=False,
        )
    )

    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    # =========================================================================
    # SHEET 1: Summary Sheet (All 92 Companies)
    # =========================================================================
    ws_summary = wb.create_sheet(title="Summary")

    # Build summary DataFrame
    summary_data = universe_df[
        ["company_id", "company_name", "broad_sector", "composite_ranking_score"]
    ].copy()
    summary_data["peer_group"] = summary_data["company_id"].map(peer_group_map)

    # Preset flags
    for p_key, p_df in all_presets.items():
        summary_data[f"in_{p_key}"] = summary_data["company_id"].isin(
            set(p_df["company_id"])
        )

    preset_flag_cols = [f"in_{k}" for k in all_presets]
    summary_data["total_presets_matched"] = summary_data[preset_flag_cols].sum(axis=1)

    # Sort summary: score descending, then preset count
    summary_data = summary_data.sort_values(
        by=["composite_ranking_score", "total_presets_matched"],
        ascending=[False, False],
        na_position="last",
    )

    summary_headers = (
        [
            "company_id",
            "company_name",
            "broad_sector",
            "peer_group",
            "composite_ranking_score",
        ]
        + preset_flag_cols
        + ["total_presets_matched"]
    )

    # Write summary headers
    for col_idx, h_name in enumerate(summary_headers, start=1):
        cell = ws_summary.cell(row=1, column=col_idx, value=h_name)
        cell.fill = PatternFill(
            start_color="1E3A8A", end_color="1E3A8A", fill_type="solid"
        )
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Write summary rows
    for row_idx, (_, r_data) in enumerate(summary_data.iterrows(), start=2):
        for col_idx, h_name in enumerate(summary_headers, start=1):
            val = r_data[h_name]
            cell = ws_summary.cell(row=row_idx, column=col_idx)
            if pd.isna(val):
                cell.value = None
            elif isinstance(val, (bool, bool)):
                cell.value = bool(val)
                cell.alignment = Alignment(horizontal="center")
            elif isinstance(val, (int, float)):
                cell.value = round(float(val), 2)
                cell.alignment = Alignment(horizontal="right")
            else:
                cell.value = str(val)
                cell.alignment = Alignment(horizontal="left")
            cell.border = thin_border

    # ColorScaleRule on composite_ranking_score column (col 5)
    score_rule = ColorScaleRule(
        start_type="min",
        start_color=COLOR_RED,
        mid_type="percentile",
        mid_value=50,
        mid_color=COLOR_YELLOW,
        end_type="max",
        end_color=COLOR_GREEN,
    )
    ws_summary.conditional_formatting.add(f"E2:E{len(summary_data) + 1}", score_rule)

    # Auto-adjust column widths for Summary sheet
    for col in ws_summary.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_summary.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # =========================================================================
    # SHEETS 2-12: 11 Peer Group Comparison Sheets
    # =========================================================================
    unique_groups = peer_groups_df["peer_group_name"].unique()

    for group_name in unique_groups:
        ws_group = wb.create_sheet(title=str(group_name)[:31])

        group_members = peer_groups_df[peer_groups_df["peer_group_name"] == group_name]
        company_ids = group_members["company_id"].tolist()
        benchmark_ids = set(
            group_members[group_members["is_benchmark"] == 1]["company_id"]
        )

        max_col = 1 + len(company_ids)

        # Row 1: Header Row with Company IDs as Columns
        c_head = ws_group.cell(row=1, column=1, value="Metric")
        _style_header_cell(c_head, is_metric_col=True)

        for c_idx, cid in enumerate(company_ids, start=2):
            cell = ws_group.cell(row=1, column=c_idx, value=cid)
            is_bench = cid in benchmark_ids
            _style_header_cell(cell, is_benchmark=is_bench, is_metric_col=False)

        # Row 2: Company Legal Name
        c_sub = ws_group.cell(row=2, column=1, value="Company Name")
        c_sub.font = Font(name="Calibri", size=9, bold=True, color="64748B")
        c_sub.alignment = Alignment(horizontal="left", vertical="center")

        for c_idx, cid in enumerate(company_ids, start=2):
            c_name = comp_name_map.get(cid, cid)
            cell = ws_group.cell(row=2, column=c_idx, value=c_name)
            cell.font = Font(name="Calibri", size=9, italic=True, color="475569")
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            cell.border = thin_border

        # Row 3: Section 1 Banner: METRIC VALUES
        _style_section_banner(
            ws_group, row_idx=3, title="FINANCIAL METRICS (LATEST FY)", max_col=max_col
        )

        # Filter percentiles for this group
        group_pp = peer_percentiles_df[
            peer_percentiles_df["peer_group_name"] == group_name
        ]

        # Rows 4 to 23: 20 Metric Values
        start_metric_row = 4
        for m_idx, (m_db_key, m_label, lower_is_better) in enumerate(
            PEER_WORKBOOK_METRICS
        ):
            current_row = start_metric_row + m_idx

            # Column A: Metric Name
            m_cell = ws_group.cell(row=current_row, column=1, value=m_label)
            m_cell.font = Font(name="Calibri", size=10, bold=True, color="1E293B")
            m_cell.alignment = Alignment(horizontal="left", vertical="center")
            m_cell.border = thin_border

            # Columns B..: Company Values
            for c_idx, cid in enumerate(company_ids, start=2):
                val_row = group_pp[
                    (group_pp["company_id"] == cid)
                    & (group_pp["metric_name"] == m_db_key)
                ]
                cell = ws_group.cell(row=current_row, column=c_idx)
                cell.border = thin_border

                if not val_row.empty and pd.notna(val_row["metric_value"].iloc[0]):
                    val = float(val_row["metric_value"].iloc[0])
                    cell.value = val
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    cell.value = None
                    cell.alignment = Alignment(horizontal="center", vertical="center")

            # Apply ColorScaleRule across this metric row
            start_col_letter = get_column_letter(2)
            end_col_letter = get_column_letter(max_col)
            row_range = f"{start_col_letter}{current_row}:{end_col_letter}{current_row}"

            # Invert colors for metrics where lower is better (D/E, PE, PB, EV/EBITDA, CapEx Intensity)
            # Green = Top performer (highest percentile / best performance)
            # Red = Bottom performer
            if lower_is_better:
                metric_rule = ColorScaleRule(
                    start_type="min",
                    start_color=COLOR_GREEN,
                    mid_type="percentile",
                    mid_value=50,
                    mid_color=COLOR_YELLOW,
                    end_type="max",
                    end_color=COLOR_RED,
                )
            else:
                metric_rule = ColorScaleRule(
                    start_type="min",
                    start_color=COLOR_RED,
                    mid_type="percentile",
                    mid_value=50,
                    mid_color=COLOR_YELLOW,
                    end_type="max",
                    end_color=COLOR_GREEN,
                )

            ws_group.conditional_formatting.add(row_range, metric_rule)

        # Row 24: Separator
        # Row 25: Section 2 Banner: BENCHMARK GAP %
        bench_label = next(iter(benchmark_ids)) if benchmark_ids else "Benchmark"
        _style_section_banner(
            ws_group,
            row_idx=25,
            title=f"BENCHMARK GAP (%) RELATIVE TO {bench_label} (BENCHMARK)",
            max_col=max_col,
        )

        # Rows 26 to 45: 20 Benchmark Gap % Rows
        start_gap_row = 26
        for m_idx, (m_db_key, m_label, _) in enumerate(PEER_WORKBOOK_METRICS):
            current_row = start_gap_row + m_idx

            # Column A: Gap Label
            g_cell = ws_group.cell(
                row=current_row, column=1, value=f"{m_label} Gap (%)"
            )
            g_cell.font = Font(name="Calibri", size=9, bold=False, color="475569")
            g_cell.alignment = Alignment(horizontal="left", vertical="center")
            g_cell.border = thin_border

            # Columns B..: Benchmark Gap Values
            for c_idx, cid in enumerate(company_ids, start=2):
                val_row = group_pp[
                    (group_pp["company_id"] == cid)
                    & (group_pp["metric_name"] == m_db_key)
                ]
                cell = ws_group.cell(row=current_row, column=c_idx)
                cell.border = thin_border

                if not val_row.empty and pd.notna(val_row["benchmark_gap_pct"].iloc[0]):
                    gap_val = float(val_row["benchmark_gap_pct"].iloc[0])
                    cell.value = gap_val
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    cell.value = None
                    cell.alignment = Alignment(horizontal="center", vertical="center")

            # Apply ColorScaleRule on gap row: Green = above benchmark, Red = below benchmark
            start_col_letter = get_column_letter(2)
            end_col_letter = get_column_letter(max_col)
            gap_range = f"{start_col_letter}{current_row}:{end_col_letter}{current_row}"

            gap_rule = ColorScaleRule(
                start_type="min",
                start_color=COLOR_RED,
                mid_type="num",
                mid_value=0.0,
                mid_color=COLOR_YELLOW,
                end_type="max",
                end_color=COLOR_GREEN,
            )
            ws_group.conditional_formatting.add(gap_range, gap_rule)

        # Set freeze panes so Metric column and Header remain stationary
        ws_group.freeze_panes = "B3"

        # Auto-adjust column widths
        ws_group.column_dimensions["A"].width = 28
        for col_idx in range(2, max_col + 1):
            col_letter = get_column_letter(col_idx)
            ws_group.column_dimensions[col_letter].width = 16

    wb.save(out_file)
    logger.info(
        "Successfully generated peer_comparison.xlsx at %s (12 sheets)", out_file
    )

    # Mirror to root directory if saved under reports/
    if out_file.name == "peer_comparison.xlsx" and out_file.parent.name == "reports":
        root_mirror = Path("peer_comparison.xlsx")
        shutil.copy2(out_file, root_mirror)
        logger.info("Mirrored workbook to root %s", root_mirror)

    return out_file


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO)
    p = generate_peer_comparison_workbook()
    wb_test = openpyxl.load_workbook(p)
    print(f"Generated {p} with {len(wb_test.sheetnames)} sheets: {wb_test.sheetnames}")
    wb_test.close()
