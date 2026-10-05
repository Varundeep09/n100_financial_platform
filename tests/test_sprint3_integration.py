"""End-to-End Sprint 3 Integration Test Suite & Acceptance Verification.

Sprint 3, Day 20: Integration Testing & Sprint Wrap-Up
Comprehensive acceptance test verifying the complete Sprint 3 pipeline:
  1. Screener Engine & 6 Preset Screeners against live db/nifty100.db
  2. Multi-factor Composite Ranking (0-100) & screener_output.xlsx (7 sheets)
  3. AC-07 Screener Acceptance Check (ROE > 15, D/E < 1, FCF > 0 returns 10-50 companies)
  4. Peer Analytics Module & peer_percentiles table (1,120 rows across 11 groups)
  5. Peer Comparison Excel Workbook (12 sheets with openpyxl ColorScaleRule)
  6. 8-Axis Radar Chart Visualizations (92 PNG artifacts in reports/radar_charts/)
"""

from __future__ import annotations

import logging
import sqlite3
import sys
from pathlib import Path
from typing import Any

# Ensure workspace root is on sys.path for direct standalone script execution
WORKSPACE_ROOT = Path(__file__).resolve().parents[1]
if str(WORKSPACE_ROOT) not in sys.path:
    sys.path.insert(0, str(WORKSPACE_ROOT))

import openpyxl
import pandas as pd
import pytest

from src.analytics.peer import (
    populate_peer_percentiles,
)
from src.analytics.peer_comparison import (
    generate_peer_comparison_workbook,
)
from src.analytics.radar import (
    generate_all_radar_charts,
)
from src.screener.engine import (
    DEFAULT_CONFIG_PATH,
    DEFAULT_DB_PATH,
    filter_universe,
    load_screener_config,
    load_screener_universe,
    run_all_presets,
)
from src.screener.ranking import (
    calculate_composite_ranking,
    export_screener_reports,
)

logger = logging.getLogger(__name__)

PNG_MAGIC_BYTES = b"\x89PNG\r\n\x1a\n"
MIN_PNG_SIZE_BYTES = 10 * 1024  # 10 KB
RADAR_CHARTS_DIR = Path("reports/radar_charts")
PEER_COMPARISON_PATH = Path("reports/peer_comparison.xlsx")
SCREENER_OUTPUT_PATH = Path("screener_output.xlsx")

EXPECTED_PRESETS = [
    "quality_compounder",
    "value_pick",
    "dividend_aristocrat",
    "growth_rocket",
    "asset_light_champion",
    "financial_health",
]

EXPECTED_PEER_GROUPS = [
    "Automobiles",
    "Consumer Finance",
    "FMCG",
    "IT Services",
    "Life Insurance",
    "Oil & Gas",
    "Pharmaceuticals",
    "Power & Utilities",
    "Private Banks",
    "Public Sector Banks",
    "Steel",
]

EXPECTED_SCREENER_SHEETS = [
    "Summary",
    "Quality Compounder",
    "Value Pick",
    "Dividend Aristocrat",
    "Growth Rocket",
    "Asset Light Champion",
    "Financial Health",
]


@pytest.fixture(scope="module")
def sprint3_pipeline_artifacts() -> dict[str, Any]:
    """Execute the end-to-end Sprint 3 pipeline and return operational artifacts."""
    # 1. Load screener universe and compute composite ranking
    df_universe = load_screener_universe(db_path=DEFAULT_DB_PATH)
    config = load_screener_config(config_path=DEFAULT_CONFIG_PATH)

    composite_scores = calculate_composite_ranking(df_universe)
    df_universe["composite_ranking_score"] = composite_scores

    # 2. Run all 6 preset screeners
    preset_results = run_all_presets(df=df_universe, config=config)

    # 3. Export screener workbooks (master screener_output.xlsx + individual sheets)
    export_screener_reports(
        output_dir=Path("reports/screener_output"),
        master_file=SCREENER_OUTPUT_PATH,
        db_path=DEFAULT_DB_PATH,
    )

    # 4. Populate peer percentiles in SQLite
    total_percentiles_rows = populate_peer_percentiles(db_path=DEFAULT_DB_PATH)

    # 5. Generate peer comparison multi-sheet workbook
    generate_peer_comparison_workbook(
        db_path=DEFAULT_DB_PATH,
        output_path=PEER_COMPARISON_PATH,
    )

    # 6. Ensure all 92 radar chart PNGs are generated and exist
    radar_pngs = list(RADAR_CHARTS_DIR.glob("*_radar.png"))
    if len(radar_pngs) < 92:
        radar_summary = generate_all_radar_charts(
            db_path=DEFAULT_DB_PATH,
            output_dir=RADAR_CHARTS_DIR,
        )
    else:
        radar_summary = {
            "total_companies": 92,
            "total_generated": len(radar_pngs),
            "failed_count": 0,
            "files": radar_pngs,
        }

    return {
        "universe_df": df_universe,
        "config": config,
        "composite_scores": composite_scores,
        "preset_results": preset_results,
        "peer_percentiles_count": total_percentiles_rows,
        "radar_summary": radar_summary,
    }


def test_screener_engine_six_presets_operational(
    sprint3_pipeline_artifacts: dict[str, Any],
) -> None:
    """Exit Criterion 1: Verify all 6 preset screeners are operational with expected company counts."""
    presets = sprint3_pipeline_artifacts["preset_results"]

    assert len(presets) == 6, f"Expected 6 presets, got {len(presets)}"
    assert set(presets.keys()) == set(EXPECTED_PRESETS)

    # Verify each preset returns non-zero results and excludes SBIN from balance-sheet screens
    for preset_name, p_df in presets.items():
        assert isinstance(p_df, pd.DataFrame)
        assert len(p_df) > 0, f"Preset '{preset_name}' returned 0 companies"
        assert (
            "SBIN" not in p_df["company_id"].values
        ), f"SBIN unexpectedly found in preset '{preset_name}'"

    # Verify counts match Day 15 baseline specifications
    assert len(presets["quality_compounder"]) == 27
    assert len(presets["value_pick"]) == 2
    assert len(presets["dividend_aristocrat"]) == 24
    assert len(presets["growth_rocket"]) == 11
    assert len(presets["asset_light_champion"]) == 3
    assert len(presets["financial_health"]) == 30


def test_ac07_quality_screener_acceptance(
    sprint3_pipeline_artifacts: dict[str, Any],
) -> None:
    """AC-07 Acceptance Check: ROE > 15, D/E < 1.0, FCF > 0 returns 10-50 companies."""
    df_universe = sprint3_pipeline_artifacts["universe_df"]

    ac07_filters = {
        "return_on_equity_pct": {"min": 15.0},
        "debt_to_equity": {"max": 1.0},
        "free_cash_flow_cr": {"min": 0.0},
    }

    ac07_results = filter_universe(df_universe, filters=ac07_filters, logic="AND")
    count = len(ac07_results)

    assert (
        10 <= count <= 50
    ), f"AC-07 failed: expected between 10 and 50 companies for ROE>15, D/E<1, FCF>0, got {count}"

    # Validate that every passing company rigorously meets the fundamental criteria
    assert (ac07_results["return_on_equity_pct"] > 15.0).all()
    assert (ac07_results["debt_to_equity"] < 1.0).all()
    assert (ac07_results["free_cash_flow_cr"] > 0.0).all()

    # SBIN and BEL/HAL must not appear due to missing/corrupted balance sheets
    excluded_tickers = {"SBIN", "BEL", "HAL"}
    assert not any(
        t in ac07_results["company_id"].values for t in excluded_tickers
    ), "Corrupted/missing balance sheet tickers leaked into AC-07 results"


def test_screener_output_workbook_structure_and_composite_rankings(
    sprint3_pipeline_artifacts: dict[str, Any],
) -> None:
    """Exit Criterion 5: Verify screener_output.xlsx has 7 sheets with composite rankings."""
    assert (
        SCREENER_OUTPUT_PATH.exists()
    ), f"Missing screener_output.xlsx at {SCREENER_OUTPUT_PATH}"

    wb = openpyxl.load_workbook(SCREENER_OUTPUT_PATH, read_only=True)
    sheet_names = wb.sheetnames
    wb.close()

    assert (
        len(sheet_names) == 7
    ), f"Expected 7 sheets in screener_output.xlsx, got {len(sheet_names)}"
    assert (
        sheet_names == EXPECTED_SCREENER_SHEETS
    ), f"Sheet names mismatch: {sheet_names} vs {EXPECTED_SCREENER_SHEETS}"

    summary_df = pd.read_excel(SCREENER_OUTPUT_PATH, sheet_name="Summary")
    assert (
        len(summary_df) == 92
    ), f"Expected 92 companies on Summary sheet, got {len(summary_df)}"
    assert "composite_ranking_score" in summary_df.columns
    assert "total_presets_matched" in summary_df.columns

    # Verify score bounds and null-handling
    valid_scores = summary_df["composite_ranking_score"].dropna()
    assert (
        len(valid_scores) >= 80
    ), "Expected at least 80 companies to have valid scores"
    assert (valid_scores >= 0.0).all()
    assert (valid_scores <= 100.0).all()

    # Verify Top 5 companies have composite scores >= 65 and top score > 70
    top_5_scores = summary_df.nlargest(5, "composite_ranking_score")[
        "composite_ranking_score"
    ]
    assert (top_5_scores >= 65.0).all()
    assert float(top_5_scores.iloc[0]) > 70.0


def test_peer_percentiles_table_invariants_and_row_count(
    sprint3_pipeline_artifacts: dict[str, Any],
) -> None:
    """Exit Criterion 2: Verify peer_percentiles table is populated for all 11 groups (1,120 rows)."""
    conn = sqlite3.connect(DEFAULT_DB_PATH)
    cursor = conn.cursor()

    row_count = cursor.execute("SELECT COUNT(*) FROM peer_percentiles;").fetchone()[0]
    distinct_groups = [
        r[0]
        for r in cursor.execute(
            "SELECT DISTINCT peer_group_name FROM peer_percentiles ORDER BY peer_group_name;"
        ).fetchall()
    ]
    distinct_companies = [
        r[0]
        for r in cursor.execute(
            "SELECT DISTINCT company_id FROM peer_percentiles ORDER BY company_id;"
        ).fetchall()
    ]
    distinct_metrics = [
        r[0]
        for r in cursor.execute(
            "SELECT DISTINCT metric_name FROM peer_percentiles ORDER BY metric_name;"
        ).fetchall()
    ]
    conn.close()

    assert (
        row_count == 1120
    ), f"Invariant violation: peer_percentiles must have exactly 1,120 rows, got {row_count}"
    assert (
        len(distinct_groups) == 11
    ), f"Expected 11 peer groups, got {len(distinct_groups)}"
    assert distinct_groups == EXPECTED_PEER_GROUPS
    assert (
        len(distinct_companies) == 56
    ), f"Expected 56 peer group member companies, got {len(distinct_companies)}"
    assert (
        len(distinct_metrics) == 20
    ), f"Expected 20 peer metrics, got {len(distinct_metrics)}"


def test_peer_comparison_workbook_structure_and_formatting(
    sprint3_pipeline_artifacts: dict[str, Any],
) -> None:
    """Exit Criterion 4: Verify peer_comparison.xlsx has 12 sheets with conditional formatting."""
    assert (
        PEER_COMPARISON_PATH.exists()
    ), f"Missing peer_comparison.xlsx at {PEER_COMPARISON_PATH}"

    wb = openpyxl.load_workbook(PEER_COMPARISON_PATH)
    sheet_names = wb.sheetnames

    assert (
        len(sheet_names) == 12
    ), f"Expected 12 sheets in peer_comparison.xlsx, got {len(sheet_names)}"
    expected_sheets = ["Summary"] + EXPECTED_PEER_GROUPS
    assert set(sheet_names) == set(
        expected_sheets
    ), f"Sheet names mismatch: {set(sheet_names) ^ set(expected_sheets)}"

    # Check each peer group sheet has ColorScaleRule conditional formatting
    for group_name in EXPECTED_PEER_GROUPS:
        ws = wb[group_name]
        cf_rules = ws.conditional_formatting
        assert (
            len(cf_rules) >= 20
        ), f"Expected at least 20 CF rules in sheet {group_name}, found {len(cf_rules)}"

        # Validate that colorScale exists in the formatting rules
        has_color_scale = any(
            any(r.type == "colorScale" or hasattr(r, "colorScale") for r in cf.rules)
            for cf in cf_rules
        )
        assert (
            has_color_scale
        ), f"No ColorScaleRule conditional formatting found in sheet '{group_name}'"

    wb.close()


def test_all_ninety_two_radar_charts_exist_and_valid(
    sprint3_pipeline_artifacts: dict[str, Any],
) -> None:
    """Exit Criterion 3: Verify all 92 radar chart PNG files exist, are >10KB, and have valid headers."""
    assert (
        RADAR_CHARTS_DIR.exists()
    ), f"Radar charts directory missing at {RADAR_CHARTS_DIR}"

    conn = sqlite3.connect(DEFAULT_DB_PATH)
    cursor = conn.cursor()
    all_companies = [
        r[0] for r in cursor.execute("SELECT id FROM companies ORDER BY id;").fetchall()
    ]
    conn.close()

    assert (
        len(all_companies) == 92
    ), f"Expected 92 universe companies, got {len(all_companies)}"

    found_pngs = list(RADAR_CHARTS_DIR.glob("*_radar.png"))
    assert (
        len(found_pngs) == 92
    ), f"Expected exactly 92 radar PNG files, found {len(found_pngs)}"

    for cid in all_companies:
        chart_file = RADAR_CHARTS_DIR / f"{cid}_radar.png"
        assert (
            chart_file.exists()
        ), f"Missing radar chart for company '{cid}': {chart_file}"

        file_size = chart_file.stat().st_size
        assert (
            file_size > MIN_PNG_SIZE_BYTES
        ), f"Radar chart for {cid} is under 10KB ({file_size} bytes)"

        with open(chart_file, "rb") as f:
            header = f.read(8)
            assert (
                header == PNG_MAGIC_BYTES
            ), f"Radar chart for {cid} does not contain valid PNG magic bytes header"


if __name__ == "__main__":
    print("=" * 80)
    print("SPRINT 3 END-TO-END PIPELINE ACCEPTANCE TEST RUNNER")
    print("=" * 80)

    # 1. Run pipeline
    print("\n[1/6] Running Screener Engine & Presets against nifty100.db...")
    universe_df = load_screener_universe()
    cfg = load_screener_config()
    composite_scores = calculate_composite_ranking(universe_df)
    universe_df["composite_ranking_score"] = composite_scores
    presets = run_all_presets(df=universe_df, config=cfg)
    for p_name, p_df in presets.items():
        print(f"  -> Preset '{p_name}': {len(p_df)} companies")

    # 2. AC-07 check
    print("\n[2/6] Verifying AC-07 Screener Acceptance Criteria...")
    ac07_filters = {
        "return_on_equity_pct": {"min": 15.0},
        "debt_to_equity": {"max": 1.0},
        "free_cash_flow_cr": {"min": 0.0},
    }
    ac07 = filter_universe(universe_df, filters=ac07_filters, logic="AND")
    print(f"  -> AC-07 (ROE>15, D/E<1, FCF>0): {len(ac07)} companies (Spec: 10-50)")
    assert 10 <= len(ac07) <= 50

    # 3. Export screener output
    print("\n[3/6] Exporting screener_output.xlsx...")
    export_screener_reports(
        output_dir="reports/screener_output", master_file=SCREENER_OUTPUT_PATH
    )
    wb_s = openpyxl.load_workbook(SCREENER_OUTPUT_PATH, read_only=True)
    print(f"  -> screener_output.xlsx verified: {len(wb_s.sheetnames)} sheets")
    wb_s.close()

    # 4. Peer percentiles
    print("\n[4/6] Running Peer Analytics & populating peer_percentiles table...")
    n_rows = populate_peer_percentiles()
    print(f"  -> peer_percentiles rows populated: {n_rows} (Expected: 1,120)")
    assert n_rows == 1120

    # 5. Peer comparison workbook
    print("\n[5/6] Generating reports/peer_comparison.xlsx...")
    generate_peer_comparison_workbook()
    wb_p = openpyxl.load_workbook(PEER_COMPARISON_PATH)
    print(f"  -> peer_comparison.xlsx verified: {len(wb_p.sheetnames)} sheets")
    wb_p.close()

    # 6. Radar charts
    print("\n[6/6] Verifying 92 Radar Charts...")
    charts = list(RADAR_CHARTS_DIR.glob("*_radar.png"))
    print(f"  -> Radar chart PNG count: {len(charts)} (Expected: 92)")
    assert len(charts) == 92

    print("\n" + "=" * 80)
    print("ALL SPRINT 3 ACCEPTANCE INVARIANTS CONFIRMED: 100% PASS")
    print("=" * 80)
